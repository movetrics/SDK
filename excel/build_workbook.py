"""Build Movetrics-Cloud.xlsx: a Start-here sheet, a Queries sheet with each .pq text, and an empty Results sheet.

Power Query queries cannot be written into a workbook by this script. The person pastes
each queries/*.pq file into Power Query once. This script only puts the text in front of them.

Run from anywhere:  python3 sdk/excel/build_workbook.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
QUERY_DIR = HERE / "queries"
OUT = HERE / "Movetrics-Cloud.xlsx"

# Order matters: BaseUrl first, because the other queries read it.
QUERY_FILES = ["BaseUrl.pq", "Results.pq", "Recordings.pq", "Athletes.pq", "Samples.pq"]

# Must match the Columns list in queries/Results.pq.
RESULT_COLUMNS = [
    "recording_id", "recorded_at", "athlete_code", "athlete_name", "movement",
    "metric_id", "label", "value", "unit", "method", "method_version",
    "engine_version", "origin",
]

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE7F0")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

START_STEPS = [
    ("Make a token.", "In Movetrics Cloud, go to Account > Access tokens > New token. Copy the token. It starts with mvt_. Keep it private. Do not type it into a cell."),
    ("Open Excel.", "Use Excel for Windows or Mac (Microsoft 365, or Excel 2021 or later). Excel on the web cannot sign in to Movetrics Cloud."),
    ("Open the query editor.", "Go to Data > Get Data > Launch Power Query Editor."),
    ("Add the first blank query.", "In the editor, go to Home > New Source > Blank Query. Then go to Home > Advanced Editor."),
    ("Paste the first query.", "On the Queries sheet, copy the text in cell B2 (BaseUrl.pq). In the Advanced Editor, delete what is there, paste the text, and click Done."),
    ("Name it.", "In Query Settings, set the name to the file name without .pq. For BaseUrl.pq, the name is BaseUrl."),
    ("Repeat for the other queries.", "Do steps 4 to 6 for Results.pq, Recordings.pq, Athletes.pq and Samples.pq. Use the same name rule. Samples is a function, so name it Samples."),
    ("Load Results to this sheet.", "Select the Results query, then Home > Close & Load To. Choose Table, then the Results sheet. Load the other tables the same way if you want them, on their own sheets."),
    ("Sign in once.", "When Excel asks, choose Basic. User name: anything, for example token. Password: your mvt_ token. Click Connect. Excel keeps the token in its own store, not in this file."),
    ("Refresh.", "Go to Data > Refresh All. New results come in. Excel reads the token from its own store, so you do not sign in again."),
    ("Read one recording's samples.", "Make a blank query. In the Advanced Editor, type: Samples(\"the-recording-id\"). Use the id from the Recordings table, between the quote marks. Then Close & Load To a sheet."),
]

STATUS_LINES = [
    "Status of this workbook:",
    "The Power Query code was checked by reading it. It has not been run in Excel yet.",
    "This script does not put the queries into the workbook. You paste them once (about two minutes).",
    "Excel on the web cannot sign in to the API, so do not use it.",
]


def write_start(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Start here"
    ws["A1"] = "Movetrics Cloud in Excel: start here"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Results and raw samples from your Movetrics Cloud workspace, refreshed with one click."

    row = 4
    for line in STATUS_LINES:
        ws.cell(row=row, column=1, value=line)
        if line.endswith(":"):
            ws.cell(row=row, column=1).font = BOLD
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Steps").font = BOLD
    row += 1
    for number, (title, text) in enumerate(START_STEPS, start=1):
        ws.cell(row=row, column=1, value=number).alignment = Alignment(vertical="top", horizontal="left")
        ws.cell(row=row, column=2, value=title).font = BOLD
        ws.cell(row=row, column=2).alignment = WRAP_TOP
        ws.cell(row=row, column=3, value=text).alignment = WRAP_TOP
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Files and links").font = BOLD
    row += 1
    ws.cell(row=row, column=2, value="Query code (paste these)")
    ws.cell(row=row, column=3, value="The Queries sheet in this workbook")
    row += 1
    ws.cell(row=row, column=2, value="Query files and source")
    ws.cell(row=row, column=3, value="https://github.com/movetrics/sdk/tree/main/excel")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 100


def write_queries(wb: Workbook) -> None:
    ws = wb.create_sheet("Queries")
    ws["A1"] = "File"
    ws["B1"] = "Power Query code (click the cell, copy the text, paste into the Advanced Editor)"
    for cell in (ws["A1"], ws["B1"]):
        cell.font = BOLD
        cell.fill = HEADER_FILL

    for row, name in enumerate(QUERY_FILES, start=2):
        text = (QUERY_DIR / name).read_text(encoding="utf-8").rstrip("\n")
        ws.cell(row=row, column=1, value=name).alignment = WRAP_TOP
        ws.cell(row=row, column=2, value=text).alignment = WRAP_TOP
        lines = text.count("\n") + 1
        ws.row_dimensions[row].height = min(409, 15 * lines)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 120


def write_results(wb: Workbook) -> None:
    ws = wb.create_sheet("Results")
    for col, name in enumerate(RESULT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = max(14, len(name) + 4)
    ws["A3"] = "Placeholder: no data yet."
    ws["A4"] = "Follow the steps on the Start here sheet. Loading the Results query here replaces this note and the headers."
    ws["A3"].font = BOLD
    ws.freeze_panes = "A2"


def main() -> None:
    wb = Workbook()
    write_start(wb)
    write_queries(wb)
    write_results(wb)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
